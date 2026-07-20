import io
import logging
import zipfile

from docx import Document

from odoo import _
from odoo.exceptions import ValidationError

from .base_extractor import BaseCvExtractor

_logger = logging.getLogger(__name__)


class DocxCvExtractor(BaseCvExtractor):
    supported_mimetypes = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    supported_extensions = (".docx",)

    def extract(self, attachment):
        docx_bytes = self._decode_attachment(attachment)
        try:
            document = Document(io.BytesIO(docx_bytes))
        except (ValueError, KeyError, TypeError, zipfile.BadZipFile) as error:
            _logger.exception("Could not read DOCX attachment %s", attachment.id)
            raise ValidationError(_("The DOCX file could not be read.")) from error

        lines = [paragraph.text.strip() for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cell_values = [cell.text.strip() for cell in row.cells]
                if any(cell_values):
                    lines.append(" | ".join(value for value in cell_values if value))
        return "\n".join(line for line in lines if line)
